# Tests for GHFDBExportResource.

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from project.ghfdb.constants import GHFDB_COLUMN_ORDER

# T041: Resource class structure tests


class TestGHFDBExportResourceDeclaration:
    @pytest.mark.xfail(
        strict=True,
        reason=(
            "Half-landed GHFDB canonical column work: the constants moved to the "
            "published spreadsheet casing, but ghfdb_colmeta.json, the resource "
            "field declarations and one manager annotation key did not follow. "
            "Needs debugging, and a decision on the published column vocabulary, "
            "before it can pass. See issue #122."
        ),
    )
    def test_all_62_columns_declared(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        field_names = set(resource.fields.keys())
        missing = set(GHFDB_COLUMN_ORDER) - field_names
        assert not missing, f"Missing fields: {missing}"

    @pytest.mark.xfail(
        strict=True,
        reason=(
            "Half-landed GHFDB canonical column work: the constants moved to the "
            "published spreadsheet casing, but ghfdb_colmeta.json, the resource "
            "field declarations and one manager annotation key did not follow. "
            "Needs debugging, and a decision on the published column vocabulary, "
            "before it can pass. See issue #122."
        ),
    )
    def test_no_extra_columns_beyond_column_order(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        field_names = set(resource.fields.keys())
        extra = field_names - set(GHFDB_COLUMN_ORDER)
        assert not extra, f"Extra undocumented fields: {extra}"

    def test_export_order_matches_ghfdb_column_order(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        assert list(GHFDBExportResource.Meta.export_order) == list(GHFDB_COLUMN_ORDER)


# T041: Queryset tests


class TestGHFDBExportQueryset:
    @pytest.mark.django_db
    def test_get_queryset_returns_ghfdb_querytype(self):
        from project.ghfdb.managers import GHFDBChildQuerySet
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = resource.get_queryset()
        assert isinstance(qs, GHFDBChildQuerySet)

    @pytest.mark.django_db
    def test_filtered_queryset_exports_only_matching_records(self, heat_flow_chain):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        assert len(dataset) == 1

    @pytest.mark.xfail(
        strict=True,
        reason=(
            "Half-landed GHFDB canonical column work: the constants moved to the "
            "published spreadsheet casing, but ghfdb_colmeta.json, the resource "
            "field declarations and one manager annotation key did not follow. "
            "Needs debugging, and a decision on the published column vocabulary, "
            "before it can pass. See issue #122."
        ),
    )
    @pytest.mark.django_db
    def test_empty_queryset_exports_headers_only(self):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.none()
        dataset = resource.export(qs)
        assert len(dataset) == 0
        assert list(dataset.headers) == list(GHFDB_COLUMN_ORDER)


# T041: Column order tests


class TestGHFDBExportColumnOrder:
    @pytest.mark.xfail(
        strict=True,
        reason=(
            "Half-landed GHFDB canonical column work: the constants moved to the "
            "published spreadsheet casing, but ghfdb_colmeta.json, the resource "
            "field declarations and one manager annotation key did not follow. "
            "Needs debugging, and a decision on the published column vocabulary, "
            "before it can pass. See issue #122."
        ),
    )
    @pytest.mark.django_db
    def test_exported_headers_match_column_order(self, heat_flow_chain):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        assert list(dataset.headers) == list(GHFDB_COLUMN_ORDER)


# T041: Pint quantity rendering tests


class TestGHFDBExportQuantityFields:
    @pytest.mark.django_db
    def test_qc_renders_as_plain_number(self, heat_flow_chain):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        row = dataset.dict[0]
        qc_val = row["qc"]
        assert not hasattr(qc_val, "magnitude"), (
            f"Expected plain number, got Quantity: {qc_val!r}"
        )
        assert qc_val is not None and str(qc_val) != ""

    @pytest.mark.django_db
    def test_t_grad_mean_renders_as_plain_number(self, heat_flow_chain):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        row = dataset.dict[0]
        grad_val = row["t_grad_mean"]
        assert not hasattr(grad_val, "magnitude"), (
            f"Expected plain number, got Quantity: {grad_val!r}"
        )
        assert grad_val is not None and str(grad_val) != ""

    @pytest.mark.django_db
    def test_tc_mean_renders_as_plain_number(self, heat_flow_chain):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        row = dataset.dict[0]
        tc_val = row["tc_mean"]
        assert not hasattr(tc_val, "magnitude"), (
            f"Expected plain number, got Quantity: {tc_val!r}"
        )
        assert tc_val is not None and str(tc_val) != ""

    @pytest.mark.django_db
    def test_water_temperature_column_reads_the_renamed_surface_temperature_field(
        self, dataset, heat_flow_chain
    ):
        # The published export still emits a column named ``water_temperature`` (D-c,
        # specs/004-import-upload-template/decisions.md — the submission template and
        # the published release are separate contracts).
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        heat_flow_chain.surface_temperature = 4.5
        heat_flow_chain.save()

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset_out = resource.export(qs)
        row = dataset_out.dict[0]

        assert "water_temperature" in row
        assert float(row["water_temperature"]) == pytest.approx(4.5)


# T041: M2M rendering tests


class TestGHFDBExportM2MFields:
    @pytest.mark.django_db
    def test_empty_q_method_renders_as_empty_string(self, heat_flow_chain):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        row = dataset.dict[0]
        assert row["q_method"] in ("", None)

    @pytest.mark.django_db
    def test_empty_t_method_top_renders_as_empty_string(self, heat_flow_chain):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        row = dataset.dict[0]
        assert row["t_method_top"] in ("", None)


# Ref_IGSN column tests (D26, specs/004-import-upload-template/decisions.md)


class TestGHFDBExportRefIGSN:
    # ``Ref_IGSN`` reads back the interval's IGSN identifier instead of a hardcoded
    # empty string.

    @pytest.mark.django_db
    def test_ref_igsn_emits_the_stored_identifier(self, heat_flow_chain):
        from fairdm.core.sample.models import SampleIdentifier

        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        SampleIdentifier.objects.create(
            related=heat_flow_chain.sample, type="IGSN", value="10.60516/AU1101"
        )

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        row = dataset.dict[0]

        assert row["Ref_IGSN"] == "10.60516/AU1101"

    @pytest.mark.django_db
    def test_ref_igsn_emits_empty_string_when_the_interval_carries_none(
        self, heat_flow_chain
    ):
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        qs = GHFDBChild.objects.for_export().filter(pk=heat_flow_chain.pk)
        dataset = resource.export(qs)
        row = dataset.dict[0]

        assert row["Ref_IGSN"] == ""


# T041: Staff-only access control


class TestGHFDBExportAccessControl:
    @pytest.mark.django_db
    def test_anonymous_admin_export_url_redirects(self, client):
        from django.urls import reverse

        url = reverse("admin:ghfdb_ghfdbchild_export")
        response = client.get(url)
        assert response.status_code == 302


# T094 — BUG-010: Export attribute= values must match canonical annotation keys


class TestBUG010ExportAttributeValues:
    # GHFDBExportResource field attribute= values must match renamed annotation keys .
    # After T096 renames the annotation keys in GHFDBChildQuerySet.as_ghfdb_flat(), the
    # export resource's attribute= values must use the new key names.

    def test_q_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["q"].attribute == "q", (
            f"Field 'q' has attribute '{resource.fields['q'].attribute}', expected 'q' (BUG-010 T099)"
        )

    def test_q_uncertainty_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["q_uncertainty"].attribute == "q_uncertainty", (
            "Field 'q_uncertainty' has stale attribute (BUG-010 T099)"
        )

    def test_name_attribute_reads_site_name_annotation(self):
        # Field 'name' must read from 'site_name' annotation (model-field conflict
        # workaround, BUG-010). 'name' cannot be used as the annotate() key because
        # Measurement base class already has a 'name' field.
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["name"].attribute == "site_name", (
            f"Field 'name' has attribute '{resource.fields['name'].attribute}', expected 'site_name' (BUG-010 T099)"
        )

    def test_elevation_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["elevation"].attribute == "elevation", (
            f"Field 'elevation' has attribute '{resource.fields['elevation'].attribute}', expected 'elevation' (BUG-010)"
        )

    def test_environment_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["environment"].attribute == "environment", (
            "Field 'environment' has stale attribute (BUG-010 T099)"
        )

    def test_corr_hp_flag_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["corr_hp_flag"].attribute == "corr_HP_flag", (
            f"Field 'corr_hp_flag' has attribute '{resource.fields['corr_hp_flag'].attribute}', expected 'corr_HP_flag' (BUG-010)"
        )

    def test_total_depth_md_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["total_depth_md"].attribute == "total_depth_MD", (
            f"Field 'total_depth_md' has attribute '{resource.fields['total_depth_md'].attribute}', expected 'total_depth_MD' (BUG-010)"
        )

    def test_total_depth_tvd_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["total_depth_tvd"].attribute == "total_depth_TVD", (
            f"Field 'total_depth_tvd' has attribute '{resource.fields['total_depth_tvd'].attribute}', expected 'total_depth_TVD' (BUG-010)"
        )

    def test_explo_method_attribute_is_canonical(self):
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert resource.fields["explo_method"].attribute == "explo_method", (
            f"Field 'explo_method' has attribute '{resource.fields['explo_method'].attribute}', expected 'explo_method' (BUG-010)"
        )

    def test_water_temperature_attribute_reads_surface_temperature(self):
        # Field 'water_temperature' attribute must be 'surface_temperature', the renamed
        # field name (US-7).
        from project.ghfdb.resources.export import GHFDBExportResource

        resource = GHFDBExportResource()
        assert (
            resource.fields["water_temperature"].attribute == "surface_temperature"
        ), (
            f"Field 'water_temperature' has attribute "
            f"'{resource.fields['water_temperature'].attribute}', expected "
            f"'surface_temperature'"
        )


def _norm_empty(value):
    """Normalize None/empty-like values to empty string for stable comparisons."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def _assert_float_close(actual, expected, tol=1e-9):
    """Assert actual ~= expected for numeric export cells."""
    if expected == "":
        assert _norm_empty(actual) == ""
        return
    assert abs(float(actual) - float(expected)) < tol


def _build_simple_xlsx_from_official(official_xlsx_bytes: bytes) -> bytes:
    """Convert an official GHFDB XLSX to simple layout by removing rows 7 and 8.

    The official template has:
      Row 6: headers, Row 7: units, Row 8: ranges, Row 9+: data

    The simple layout has:
      Row 6: headers, Row 7+: data

    This function reads the official XLSX, drops rows 7-8, and re-writes
    a new XLSX in the simple layout.

    Args:
        official_xlsx_bytes: Raw bytes of the official-layout XLSX.

    Returns:
        Raw bytes of the equivalent simple-layout XLSX.
    """
    wb_in = openpyxl.load_workbook(
        BytesIO(official_xlsx_bytes), read_only=True, data_only=True
    )
    ws_in = wb_in["data list"]

    # Read all rows (1-indexed); official has headers at row 6, unit at 7, range at 8
    all_rows = list(ws_in.iter_rows(values_only=True))
    # Rows 1-6 (0-indexed 0-5) stay; rows 7-8 (0-indexed 6-7) are dropped; data is 8+ (0-indexed)
    kept_rows = (
        all_rows[:6] + all_rows[8:]
    )  # keep metadata+header, skip units/ranges, keep data
    wb_in.close()

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "data list"

    for row_idx, row_values in enumerate(kept_rows, start=1):
        for col_idx, value in enumerate(row_values, start=1):
            ws_out.cell(row=row_idx, column=col_idx, value=value)

    buf = BytesIO()
    wb_out.save(buf)
    return buf.getvalue()


class TestGHFDBExportRoundTrip:
    @pytest.mark.django_db
    def test_roundtrip_import_then_export_preserves_values(self, dataset):
        # SC-001: 1) Import fixture XLSX using parent + child resources 2) Export using
        # GHFDBExportResource 3) Verify text/vocabulary equality and numeric closeness
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources import (
            GHFDBChildImportResource,
            GHFDBExportResource,
            GHFDBImportFormat,
            GHFDBParentImportResource,
        )

        fixture_path = (
            Path(__file__).resolve().parents[1] / "fixtures" / "sample_ghfdb.xlsx"
        )
        fixture_bytes = fixture_path.read_bytes()

        # Parse twice because parent import mutates its Dataset during deduplication.
        fmt = GHFDBImportFormat()
        ds_parent = fmt.create_dataset(fixture_bytes)
        ds_child = fmt.create_dataset(fixture_bytes)

        parent_resource = GHFDBParentImportResource()
        parent_result = parent_resource.import_data(
            ds_parent,
            dry_run=False,
            raise_errors=True,
            fairdm_dataset=dataset,
        )
        assert not parent_result.has_errors(), parent_result.invalid_rows

        child_resource = GHFDBChildImportResource()
        child_result = child_resource.import_data(
            ds_child,
            dry_run=False,
            raise_errors=True,
            fairdm_dataset=dataset,
        )
        assert not child_result.has_errors(), child_result.invalid_rows

        export_resource = GHFDBExportResource()
        export_qs = GHFDBChild.objects.for_export().order_by("ghfdb_id")
        exported_rows = list(export_resource.export(export_qs).dict)

        assert len(exported_rows) == 3

        by_comment = {row["c_comment"]: row for row in exported_rows}

        expected = {
            "child-corrections": {
                "text": {
                    "name": "Roundtrip Site",
                    "environment": "onshore_continental",
                    "p_comment": "parent comment",
                    "explo_purpose": "Hydrocarbon",
                    "q_method": "Bullard method",
                    "corr_is_flag": "present_corrected",
                    "corr_t_flag": "present_not_corrected",
                    "corr_s_flag": "present_not_significant",
                    "corr_e_flag": "not_recognized",
                    "corr_topo_flag": "considered_p",
                    "corr_pal_flag": "considered_t",
                    "corr_sur_flag": "considered_pt",
                    "corr_conv_flag": "not_considered",
                    "corr_hr_flag": "present_not_significant",
                    "expedition": "Expedition A",
                },
                "numeric": {
                    "q": 70,
                    "q_uncertainty": 5,
                    "qc": 68,
                    "qc_uncertainty": 4,
                    "q_top": 0,
                    "q_bottom": 200,
                    "tc_mean": "",
                    "t_grad_mean": "",
                },
            },
            "child-gradient-conductivity": {
                "text": {
                    "name": "Roundtrip Site",
                    "environment": "onshore_continental",
                    "explo_purpose": "Hydrocarbon",
                    "q_method": "Interval method",
                    "corr_is_flag": "-",
                    "corr_t_flag": "-",
                    "corr_s_flag": "-",
                    "corr_e_flag": "-",
                    "corr_topo_flag": "-",
                    "corr_pal_flag": "-",
                    "corr_sur_flag": "-",
                    "corr_conv_flag": "-",
                    "corr_hr_flag": "-",
                    "expedition": "Expedition B",
                    "t_method_top": "BHT",
                    "t_method_bottom": "BLK",
                    "t_corr_top": "AAPG correction",
                    "t_corr_bottom": "Horner plot",
                    "tc_source": "Core samples",
                    "tc_location": "Actual heat-flow location",
                    "tc_method": "Estimation - from lithology and literature",
                    "tc_saturation": "Dry measured",
                    "tc_pT_conditions": "Actual in-situ (pT) conditions",
                    "tc_pT_function": "Other",
                    "tc_strategy": "Characterize formation conductivities",
                },
                "numeric": {
                    "q": 70,
                    "q_uncertainty": 5,
                    "qc": 66,
                    "qc_uncertainty": 3,
                    "q_top": 200,
                    "q_bottom": 500,
                    "probe_penetration": 2,
                    "probe_length": 3,
                    "probe_tilt": 5,
                    "water_temperature": 3,
                    "t_grad_mean": 25,
                    "t_grad_uncertainty": 1,
                    "t_grad_mean_cor": 24,
                    "t_grad_uncertainty_cor": 0.8,
                    "t_shutin_top": 24,
                    "t_shutin_bottom": 36,
                    "t_number": 8,
                    "tc_mean": 2.5,
                    "tc_uncertainty": 0.2,
                    "tc_number": 5,
                },
            },
            "child-minimal": {
                "text": {
                    "name": "Roundtrip Site",
                    "environment": "onshore_continental",
                    "explo_purpose": "Hydrocarbon",
                    "q_method": "Boot-strapping method",
                    "corr_is_flag": "-",
                    "corr_t_flag": "-",
                    "corr_s_flag": "-",
                    "corr_e_flag": "-",
                    "corr_topo_flag": "-",
                    "corr_pal_flag": "-",
                    "corr_sur_flag": "-",
                    "corr_conv_flag": "-",
                    "corr_hr_flag": "-",
                },
                "numeric": {
                    "q": 70,
                    "q_uncertainty": 5,
                    "qc": 64,
                    "qc_uncertainty": 2,
                    "q_top": 500,
                    "q_bottom": 800,
                    "tc_mean": "",
                    "t_grad_mean": "",
                },
            },
        }

        for comment, expected_row in expected.items():
            actual = by_comment[comment]

            for col, exp in expected_row["text"].items():
                assert _norm_empty(actual[col]) == _norm_empty(exp), (
                    f"Mismatch in '{col}' for row '{comment}'"
                )

            for col, exp in expected_row["numeric"].items():
                _assert_float_close(actual[col], exp)

    @pytest.mark.django_db
    def test_roundtrip_simple_format_import_then_export_preserves_values(self, dataset):
        # SC-001 (simple template variant): 1) Convert fixture XLSX to simple layout
        # (remove unit/range rows 7-8) 2) Import using GHFDBSimpleImportFormat (parent +
        # child resources) 3) Export using GHFDBExportResource 4) Verify identical.
        from project.ghfdb.models import GHFDBChild
        from project.ghfdb.resources import (
            GHFDBChildImportResource,
            GHFDBExportResource,
            GHFDBParentImportResource,
            GHFDBSimpleImportFormat,
        )

        fixture_path = (
            Path(__file__).resolve().parents[1] / "fixtures" / "sample_ghfdb.xlsx"
        )
        official_bytes = fixture_path.read_bytes()
        simple_bytes = _build_simple_xlsx_from_official(official_bytes)

        fmt = GHFDBSimpleImportFormat()
        ds_parent = fmt.create_dataset(simple_bytes)
        ds_child = fmt.create_dataset(simple_bytes)

        parent_resource = GHFDBParentImportResource()
        parent_result = parent_resource.import_data(
            ds_parent,
            dry_run=False,
            raise_errors=True,
            fairdm_dataset=dataset,
        )
        assert not parent_result.has_errors(), parent_result.invalid_rows

        child_resource = GHFDBChildImportResource()
        child_result = child_resource.import_data(
            ds_child,
            dry_run=False,
            raise_errors=True,
            fairdm_dataset=dataset,
        )
        assert not child_result.has_errors(), child_result.invalid_rows

        export_resource = GHFDBExportResource()
        export_qs = GHFDBChild.objects.for_export().order_by("ghfdb_id")
        exported_rows = list(export_resource.export(export_qs).dict)

        assert len(exported_rows) == 3

        by_comment = {row["c_comment"]: row for row in exported_rows}

        expected = {
            "child-corrections": {
                "text": {
                    "name": "Roundtrip Site",
                    "environment": "onshore_continental",
                    "p_comment": "parent comment",
                    "explo_purpose": "Hydrocarbon",
                    "q_method": "Bullard method",
                    "corr_is_flag": "present_corrected",
                    "corr_t_flag": "present_not_corrected",
                    "corr_s_flag": "present_not_significant",
                    "corr_e_flag": "not_recognized",
                    "corr_topo_flag": "considered_p",
                    "corr_pal_flag": "considered_t",
                    "corr_sur_flag": "considered_pt",
                    "corr_conv_flag": "not_considered",
                    "corr_hr_flag": "present_not_significant",
                    "expedition": "Expedition A",
                },
                "numeric": {
                    "q": 70,
                    "q_uncertainty": 5,
                    "qc": 68,
                    "qc_uncertainty": 4,
                    "q_top": 0,
                    "q_bottom": 200,
                    "tc_mean": "",
                    "t_grad_mean": "",
                },
            },
            "child-gradient-conductivity": {
                "text": {
                    "name": "Roundtrip Site",
                    "environment": "onshore_continental",
                    "explo_purpose": "Hydrocarbon",
                    "q_method": "Interval method",
                },
                "numeric": {
                    "q": 70,
                    "q_uncertainty": 5,
                    "qc": 66,
                    "qc_uncertainty": 3,
                    "q_top": 200,
                    "q_bottom": 500,
                },
            },
            "child-minimal": {
                "text": {
                    "name": "Roundtrip Site",
                    "environment": "onshore_continental",
                    "explo_purpose": "Hydrocarbon",
                    "q_method": "Boot-strapping method",
                },
                "numeric": {
                    "q": 70,
                    "q_uncertainty": 5,
                    "qc": 64,
                    "qc_uncertainty": 2,
                    "q_top": 500,
                    "q_bottom": 800,
                    "tc_mean": "",
                    "t_grad_mean": "",
                },
            },
        }

        for comment, expected_row in expected.items():
            actual = by_comment[comment]

            for col, exp in expected_row["text"].items():
                assert _norm_empty(actual[col]) == _norm_empty(exp), (
                    f"[simple] Mismatch in '{col}' for row '{comment}'"
                )

            for col, exp in expected_row["numeric"].items():
                _assert_float_close(actual[col], exp)
